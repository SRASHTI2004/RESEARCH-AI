"""ATS-friendly DOCX and PDF export.

ATS rules followed: one column, no tables/text boxes/images/icons, the
standard section headings parsers look for (SUMMARY, SKILLS, EXPERIENCE,
PROJECTS, EDUCATION), plain bullets, a standard font, real (selectable)
text, and contact details in the body rather than a header/footer.
"""

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

from app.core.resume import Resume


def _contact_line(resume: Resume) -> str:
    c = resume.contact
    return " | ".join(x for x in [c.email, c.phone, c.location, *c.links.values()] if x)


def _dates(start: str, end: str) -> str:
    return f"{start} – {end}" if start or end else ""


def _project_heading(name: str, tech: list[str]) -> str:
    return f"{name} | {', '.join(tech)}" if tech else name


def _sections(resume: Resume) -> list[tuple[str, list[tuple[str, str, list[str]]]]]:
    """(heading, [(line, right-side text, bullets)]) — shared by both formats."""
    sections: list[tuple[str, list[tuple[str, str, list[str]]]]] = []
    if resume.summary:
        sections.append(("SUMMARY", [(resume.summary, "", [])]))
    if resume.skills:
        sections.append(("SKILLS", [(f"{g.group}: {', '.join(g.items)}", "", []) for g in resume.skills]))
    if resume.experience:
        sections.append(
            (
                "EXPERIENCE",
                [
                    (
                        f"{e.title} — {e.org}{f', {e.location}' if e.location else ''}",
                        _dates(e.start, e.end),
                        e.bullets,
                    )
                    for e in resume.experience
                ],
            )
        )
    if resume.projects:
        sections.append(
            (
                "PROJECTS",
                [(_project_heading(p.name, p.tech), p.link, p.bullets) for p in resume.projects],
            )
        )
    if resume.education:
        sections.append(
            (
                "EDUCATION",
                [
                    (
                        f"{ed.degree} — {ed.school}{f', {ed.score}' if ed.score else ''}",
                        _dates(ed.start, ed.end),
                        [],
                    )
                    for ed in resume.education
                ],
            )
        )
    if resume.achievements:
        sections.append(("ACHIEVEMENTS", [("", "", resume.achievements)]))
    return sections


def to_docx(resume: Resume) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.6)
        section.left_margin = section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(resume.contact.name)
    run.bold = True
    run.font.size = Pt(16)
    contact = doc.add_paragraph(_contact_line(resume))
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for heading, entries in _sections(resume):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        hr = h.add_run(heading)
        hr.bold = True
        hr.font.size = Pt(11.5)
        for line, right, bullets in entries:
            if line:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                lr = p.add_run(line)
                lr.bold = heading not in ("SUMMARY", "SKILLS")
                if right:
                    p.add_run(f"  ({right})")
            for bullet in bullets:
                bp = doc.add_paragraph(bullet, style="List Bullet")
                bp.paragraph_format.space_after = Pt(0)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# fpdf2's built-in fonts are Latin-1 only (embedding a TTF would tie the
# export to fonts present on one OS) — map common typography to ASCII.
_PDF_REPLACEMENTS = {
    "–": "-", "—": "-", "‘": "'", "’": "'", "“": '"', "”": '"', "•": "-", "…": "...",
    "₹": "Rs.", "→": "->", "✓": "", " ": " ",
}  # fmt: skip


def _latin1(text: str) -> str:
    for src, dst in _PDF_REPLACEMENTS.items():
        text = text.replace(src, dst)
    return text.encode("latin-1", "replace").decode("latin-1")


def to_pdf(resume: Resume) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_margins(15, 14, 15)
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(width, 8, _latin1(resume.contact.name), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9.5)
    pdf.multi_cell(width, 5, _latin1(_contact_line(resume)), align="C", new_x="LMARGIN", new_y="NEXT")

    for heading, entries in _sections(resume):
        pdf.ln(2.5)
        pdf.set_font("Helvetica", "B", 11.5)
        pdf.cell(width, 6, heading, new_x="LMARGIN", new_y="NEXT")
        y = pdf.get_y()
        pdf.line(pdf.l_margin, y, pdf.l_margin + width, y)
        pdf.ln(1)
        for line, right, bullets in entries:
            if line:
                bold = heading not in ("SUMMARY", "SKILLS")
                pdf.set_font("Helvetica", "B" if bold else "", 10)
                text = f"{line}  ({right})" if right else line
                pdf.multi_cell(width, 5, _latin1(text), new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", "", 10)
            for bullet in bullets:
                pdf.set_x(pdf.l_margin + 3)
                pdf.multi_cell(width - 3, 5, _latin1(f"- {bullet}"), new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())


def export_filename(resume: Resume, company: str, extension: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", f"{resume.contact.name}_{company}").strip("_")
    return f"{slug or 'resume'}_resume.{extension}"
