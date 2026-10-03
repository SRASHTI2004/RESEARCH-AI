import io
import json

import pytest
from docx import Document
from pypdf import PdfReader

from app.core.llm import LLMError
from app.core.resume import load_master_resume
from app.services import resume_export, resume_tailor
from tests.factories import make_job

JOB_DESCRIPTION = (
    "We're hiring a backend developer to build RESTful APIs with Python and FastAPI on PostgreSQL. "
    "Experience with Docker and Kubernetes is a plus. You will write tests with pytest."
)


@pytest.fixture
def master():
    return load_master_resume("data/master_resume.example.yaml")


@pytest.fixture
def job():
    return make_job(title="Backend Developer", company="Acme", description=JOB_DESCRIPTION)


def _llm_returning(payload: dict):
    return lambda prompt, **kw: "```json\n" + json.dumps(payload) + "\n```"


# --- the no-invention validator -------------------------------------------------

ORIGINAL = "Built REST API endpoints in FastAPI for an internal reporting tool used by the operations team."
ITEM_TEXT = f"Software Engineering Intern Example Company\nPython\nFastAPI\nPostgreSQL\n{ORIGINAL}"


@pytest.mark.parametrize(
    ("rewrite", "reason_fragment"),
    [
        ("Built REST API endpoints in FastAPI serving 10,000 users for the operations team.", "numbers"),
        ("Built REST API endpoints in FastAPI and Kubernetes for the operations team.", "technologies"),
        ("Built REST API endpoints in FastAPI for Google's internal reporting tool.", "names"),
        ("Led a team of five building REST APIs in FastAPI for operations.", "numbers"),
        (ORIGINAL + " " + ORIGINAL, "longer"),
        ("   ", "empty"),
    ],
)
def test_rewrites_that_invent_are_rejected(rewrite, reason_fragment):
    reason = resume_tailor.check_rewrite(ORIGINAL, rewrite, allowed_text=ITEM_TEXT, skills=["Python"])
    assert reason is not None and reason_fragment in reason


@pytest.mark.parametrize(
    "rewrite",
    [
        "Developed RESTful APIs in FastAPI for an internal reporting tool used by the operations team.",
        "Built FastAPI REST API endpoints backed by PostgreSQL for an internal operations reporting tool.",
        "Designed Python REST APIs with FastAPI powering an internal reporting tool for operations.",
    ],
)
def test_honest_rewordings_are_accepted(rewrite):
    """Synonyms ("RESTful"), the item's own tech list (PostgreSQL), and job
    vocabulary for things already there are fine."""
    assert resume_tailor.check_rewrite(ORIGINAL, rewrite, allowed_text=ITEM_TEXT, skills=["Python"]) is None


def test_numbers_already_in_the_bullet_may_stay():
    original = "Cut page load time from 4s to 1.5s by adding caching."
    assert (
        resume_tailor.check_rewrite(
            original, "Reduced page load from 4s to 1.5s with caching.", allowed_text=original, skills=[]
        )
        is None
    )


# --- ordering ----------------------------------------------------------------


def test_reorder_moves_relevant_items_first_and_never_drops_anything(master):
    job_terms = {"docker", "github actions", "celery"}
    tailored = resume_tailor.reorder(master, job_terms)

    tools = next(g for g in tailored.skills if g.group == "Tools")
    assert tools.items[0] == "Docker"
    # Same content, just re-ordered.
    for before, after in zip(master.skills, tailored.skills, strict=True):
        assert sorted(before.items) == sorted(after.items)
    for p_before, p_after in zip(
        master.projects, sorted(tailored.projects, key=lambda p: p.id), strict=False
    ):
        assert sorted(p_before.bullets) == sorted(p_after.bullets)
    assert [e.id for e in tailored.experience] == [e.id for e in master.experience]  # chronology kept
    assert master.skills[2].items[0] == "Git"  # master untouched (deep copy)


# --- end-to-end tailoring ------------------------------------------------------------


def test_tailor_applies_valid_rewrites_and_rejects_invented_ones(master, job, monkeypatch):
    monkeypatch.setattr(
        "app.core.llm.invoke_llm",
        _llm_returning(
            {
                "summary": "Backend-focused full-stack developer with a 3-month internship building FastAPI services and React apps with Python and TypeScript.",
                "bullets": {
                    "exp:internship:0": "Developed RESTful APIs in FastAPI for an internal reporting tool used by the operations team.",
                    "exp:internship:1": "Wrote pytest tests for 50 services and deployed them on Kubernetes.",
                    "proj:nope:0": "Unknown id is ignored.",
                },
            }
        ),
    )
    result = resume_tailor.tailor(master, job)

    assert result.used_llm
    bullets = result.resume.experience[0].bullets
    assert (
        "Developed RESTful APIs in FastAPI for an internal reporting tool used by the operations team."
        in bullets
    )
    assert not any("Kubernetes" in b or "50" in b for b in bullets)
    assert "Wrote pytest tests for existing services and fixed bugs found during code review." in bullets
    assert any("Kept your original wording" in w and "numbers" in w for w in result.warnings)
    assert result.resume.summary.startswith("Backend-focused")

    assert "fastapi" in result.keywords_matched and "pytest" in result.keywords_matched
    assert "kubernetes" in result.keywords_missing  # a gap, reported — never added
    assert "Kubernetes" not in "\n".join(resume_tailor.render_text(result.resume))

    ops = {line.op for line in result.diff}
    assert {"-", "+", " "} <= ops


def test_invented_summary_is_rejected(master, job, monkeypatch):
    monkeypatch.setattr(
        "app.core.llm.invoke_llm",
        _llm_returning({"summary": "Senior engineer with 5 years of Kubernetes experience.", "bullets": {}}),
    )
    result = resume_tailor.tailor(master, job)
    assert result.resume.summary == master.summary
    assert any("original summary" in w for w in result.warnings)


def test_tailoring_never_adds_or_removes_items(master, job, monkeypatch):
    monkeypatch.setattr("app.core.llm.invoke_llm", _llm_returning({"bullets": {}, "summary": ""}))
    tailored = resume_tailor.tailor(master, job).resume

    assert {e.id for e in tailored.experience} == {e.id for e in master.experience}
    assert {p.id for p in tailored.projects} == {p.id for p in master.projects}
    assert sorted(i for g in tailored.skills for i in g.items) == sorted(
        i for g in master.skills for i in g.items
    )
    assert tailored.education == master.education
    assert tailored.contact == master.contact


def test_llm_outage_falls_back_to_reorder_only(master, job, monkeypatch):
    def _down(prompt, **kw):
        raise LLMError("All configured LLM providers failed")

    monkeypatch.setattr("app.core.llm.invoke_llm", _down)
    result = resume_tailor.tailor(master, job)
    assert not result.used_llm
    assert any("reordered for the job but not reworded" in w for w in result.warnings)
    assert result.resume.summary == master.summary


def test_prompt_keys_every_bullet_and_states_the_rules(master, job):
    prompt = resume_tailor.build_prompt(master, job, ["python", "fastapi"])
    assert '"exp:internship:0"' in prompt and '"proj:researchai:2"' in prompt
    assert "Never add a fact" in prompt


# --- export ------------------------------------------------------------------------


def test_docx_export_is_single_column_with_standard_headings(master):
    doc = Document(io.BytesIO(resume_export.to_docx(master)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert doc.tables == []  # no tables — ATS parsers often scramble them
    for heading in ("SUMMARY", "SKILLS", "EXPERIENCE", "PROJECTS", "EDUCATION"):
        assert heading in text
    assert "Your Name" in text
    assert "Built REST API endpoints in FastAPI" in text


def test_pdf_export_has_real_selectable_text(master):
    data = resume_export.to_pdf(master)
    assert data.startswith(b"%PDF")
    text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)
    assert "EXPERIENCE" in text and "Software Engineering Intern" in text
    # Unicode punctuation is mapped instead of crashing the Latin-1 core font.
    assert "ResearchAI - Company Research Brief" in text


def test_export_filename():
    master = load_master_resume("data/master_resume.example.yaml")
    assert resume_export.export_filename(master, "Acme Corp.", "pdf") == "Your_Name_Acme_Corp_resume.pdf"
